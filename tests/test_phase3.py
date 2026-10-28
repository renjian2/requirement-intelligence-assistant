"""Phase 3/4: issues, stories, review, traceability, exports, evaluation, parallel calls."""
from __future__ import annotations

import csv
import io
from dataclasses import replace

import httpx
import pytest
from docx import Document
from openpyxl import load_workbook

from ria.exporters.bundle import load_bundle
from ria.exporters.docx_brd import build_brd, build_review_report
from ria.exporters.jira_csv import build_jira_csv
from ria.exporters.xlsx import build_workbook
from ria.models.analysis import RequirementRecord, StatementRecord
from ria.models.enums import (
    AnalysisMode, DocumentType, EvidenceStatus, FindingType, InputMode, ItemType, RequirementType, ReviewStatus,
    RunStatus, StatementType,
)
from ria.models.segments import DocumentStats, SourceDocument
from ria.pipeline.issues import IssueContext, detect_value_conflicts, run_rules
from ria.pipeline.orchestrator import AnalysisOrchestrator, RunOptions
from ria.providers.grok import GrokProvider
from ria.providers.mock import MockProvider
from ria.services.review import ReviewError, ReviewService
from ria.services.traceability import clarification_questions, compare_meeting_vs_brd, traceability
from ria.validation.values import find_quantities, numbers_in
from tests.test_phase2_pipeline import FakeLLM


# ------------------------------------------------------------------ helpers
def _confirm_all(service, repo, sid):
    service.load_samples(sid)
    for d in repo.list_documents(sid):
        service.confirm_extraction(sid, d.document_id)


@pytest.fixture
def demo_run(service, repo, session_id, settings):
    _confirm_all(service, repo, session_id)
    return AnalysisOrchestrator(repo, settings, MockProvider()).run(session_id, RunOptions(InputMode.COMPARE))


def _doc(doc_id, dtype):
    return SourceDocument(document_id=doc_id, session_id="s", file_name=f"{doc_id}.txt", extension=".txt",
                          document_type=dtype, sha256="x" * 64, size_bytes=1, stats=DocumentStats())


def _req(rid, text, doc, group="Approvals", status=EvidenceStatus.STATED):
    return RequirementRecord(requirement_id=rid, run_id="r", group=group, title=text[:30], description=text,
                             requirement_type=RequirementType.BUSINESS, evidence_status=status,
                             source_statement_ids=[], segment_ids=[], document_ids=[doc], evidence_verified=True)


def _ctx(reqs, docs):
    return IssueContext(run_id="r", docs={d.document_id: d for d in docs}, statements={}, requirements=reqs,
                        understandings={}, project_glossary=[], accepted_corrections=[])


# ------------------------------------------------------------------ values
@pytest.mark.parametrize("text,expected", [
    ("submit within 45 days, not 30 days", [("duration", 45.0)]),
    ("claims above twenty-five thousand", [("amount", 25000.0)]),
    ("approval above INR 1,000 within 24 hours", [("amount", 1000.0), ("duration", 1.0)]),
    ("a fee of 5% on amounts over 25k", [("percent", 5.0), ("amount", 25000.0)]),
    ("respond in 2 weeks", [("duration", 14.0)]),
])
def test_find_quantities(text, expected):
    assert [(q.kind, q.value) for q in find_quantities(text)] == expected


def test_numbers_in_normalises_words_and_commas():
    assert numbers_in("INR 25,000 or twenty five thousand or 25k") == {"25000"}


# ------------------------------------------------------------------ value conflicts
def test_value_conflict_across_documents_only():
    docs = [_doc("D1", DocumentType.MEETING_TRANSCRIPT), _doc("D2", DocumentType.BRD)]
    reqs = [_req("R1", "Claims above 500 need manager approval", "D1"),
            _req("R2", "Claims above INR 1,000 require Line Manager approval", "D2")]
    [f] = detect_value_conflicts(_ctx(reqs, docs))
    assert f.finding_type == FindingType.CONTRADICTION and f.severity.value == "HIGH"
    assert {v.value for v in f.values} == {"500 (no currency)", "INR 1,000"}


def test_no_conflict_inside_one_document_or_for_different_roles_or_currencies():
    docs = [_doc("D1", DocumentType.BUSINESS_POLICY), _doc("D2", DocumentType.BRD)]
    same_doc = [_req("R1", "Meal expenses reimbursed up to INR 1,500 per day in metro cities", "D1"),
                _req("R2", "Meal expenses reimbursed up to INR 1,000 per day in other cities", "D1")]
    assert detect_value_conflicts(_ctx(same_doc, docs)) == []
    roles = [_req("R1", "Claims above INR 1,000 require manager approval", "D1"),
             _req("R2", "Claims above INR 25,000 require finance approval", "D2")]
    assert detect_value_conflicts(_ctx(roles, docs)) == []
    cur = [_req("R1", "Claims above INR 1,000 need approval", "D1"), _req("R2", "Claims above USD 50 need approval", "D2")]
    assert detect_value_conflicts(_ctx(cur, docs)) == []


def test_demo_run_rules_find_planted_issues(demo_run, repo):
    findings = repo.get_findings(demo_run.run_id)
    types = {f.finding_type for f in findings}
    assert {FindingType.CONTRADICTION, FindingType.AMBIGUITY, FindingType.GAP, FindingType.UNDEFINED_TERM,
            FindingType.SCOPE_ISSUE, FindingType.POSSIBLE_TRANSCRIPTION_ERROR} <= types
    days = next(f for f in findings if f.finding_type == FindingType.CONTRADICTION and "time limit" in f.title)
    assert sorted(v.value for v in days.values) == ["45 days", "60 days", "90 days"]
    reqs = {r.requirement_id: r for r in repo.get_requirements(demo_run.run_id)}
    assert all(reqs[r].evidence_status == EvidenceStatus.CONFLICTING and reqs[r].confidence <= 45
               for r in days.requirement_ids)
    assert [f.finding_id for f in findings] == [f"FND-{i:03d}" for i in range(1, len(findings) + 1)]
    assert not any("25,000" in f.title for f in findings)   # spoken 25k matches INR 25,000 elsewhere


def test_demo_stories_skip_proposed_and_link_requirements(demo_run, repo):
    stories = repo.get_stories(demo_run.run_id)
    reqs = {r.requirement_id: r for r in repo.get_requirements(demo_run.run_id)}
    linked = {rid for s in stories for rid in s.requirement_ids}
    assert all(reqs[r].evidence_status not in (EvidenceStatus.PROPOSED, EvidenceStatus.UNSUPPORTED) for r in linked)
    assert not any("sms" in s.story_text.lower() for s in stories)
    scenarios = repo.get_scenarios(demo_run.run_id)
    assert all(t.story_id in {s.story_id for s in stories} for t in scenarios)
    assert stories[0].acceptance_criteria[0].ac_id == "AC-001-01"


# ------------------------------------------------------------------ AI output validation (fake Grok)
def _grok(fake):
    return GrokProvider(base_url="https://api.x.ai/v1", model="grok-test", api_key="k",
                        transport=httpx.MockTransport(fake), sleep=lambda s: None)


def test_ai_findings_and_stories_are_validated(service, repo, session_id, settings):
    _confirm_all(service, repo, session_id)
    run = AnalysisOrchestrator(repo, settings, _grok(FakeLLM())).run(
        session_id, RunOptions(InputMode.COMPARE, AnalysisMode.DEEP, max_tests_per_story=3))
    assert run.status == RunStatus.COMPLETED, run.summary.errors
    findings = repo.get_findings(run.run_id)
    ai = [f for f in findings if f.detected_by.value == "LLM"]
    assert ai and all(f.is_ai_suggestion for f in ai)
    assert not any("REQ-999" in f.requirement_ids for f in findings)
    assert not any(f.title == "Needs two" for f in findings)        # contradiction with one requirement dropped
    story = repo.get_stories(run.run_id)[0]
    assert "REQ-999" not in story.requirement_ids
    assert story.actor == "[TBC: actor]" and any("Auditor" in n for n in story.validation_notes)
    assert story.acceptance_criteria[0].unverified_values == ["7"]   # '7 days' is not in the sources
    tests = [t for t in repo.get_scenarios(run.run_id) if t.story_id == story.story_id]
    assert [t.priority.value for t in tests] == ["HIGH", "HIGH", "MEDIUM"]   # capped at 3, important first
    assert any("dropped" in n for n in story.validation_notes)


def test_parallel_calls_give_same_result(service, repo, session_id, settings):
    _confirm_all(service, repo, session_id)
    seq = AnalysisOrchestrator(repo, settings, _grok(FakeLLM())).run(session_id, RunOptions(InputMode.COMPARE, use_cache=False))
    par = AnalysisOrchestrator(repo, replace(settings, max_parallel_calls=4), _grok(FakeLLM())).run(
        session_id, RunOptions(InputMode.COMPARE, use_cache=False))
    assert par.status == RunStatus.COMPLETED
    key = lambda run: [(r.requirement_id, r.description) for r in repo.get_requirements(run.run_id)]  # noqa: E731
    assert key(seq) == key(par)
    assert par.summary.calls == seq.summary.calls


# ------------------------------------------------------------------ review
def test_review_actions_and_audit(demo_run, repo, session_id):
    svc = ReviewService(repo)
    rid = demo_run.run_id
    svc.apply(session_id=session_id, run_id=rid, item_type=ItemType.REQUIREMENT, item_id="REQ-006", action="approve",
              reviewer="Nikhil")
    r = svc.apply(session_id=session_id, run_id=rid, item_type=ItemType.REQUIREMENT, item_id="REQ-005", action="approve",
                  reviewer="Nikhil", edits={"description": "Approvals happen within [TBC: SLA]"})
    assert r.review_status == ReviewStatus.EDITED
    with pytest.raises(ReviewError):
        svc.apply(session_id=session_id, run_id=rid, item_type=ItemType.REQUIREMENT, item_id="REQ-005",
                  action="approve", edits={"evidence_status": "STATED"})     # evidence cannot be edited
    with pytest.raises(ReviewError):
        svc.apply(session_id=session_id, run_id=rid, item_type=ItemType.REQUIREMENT, item_id="REQ-005", action="resolve")
    f = svc.apply(session_id=session_id, run_id=rid, item_type=ItemType.FINDING, item_id="FND-001", action="resolve",
                  note="Finance confirmed 45 days")
    assert f.review_status == ReviewStatus.RESOLVED and f.resolution == "Finance confirmed 45 days"
    audit = repo.get_audit(rid)
    assert {a["action"] for a in audit} >= {"approve", "edit", "resolve"} and audit[0]["reviewer"] == "Nikhil"


def test_merge_requirements(demo_run, repo, session_id):
    svc = ReviewService(repo)
    keep = svc.merge_requirements(session_id=session_id, run_id=demo_run.run_id, keep_id="REQ-009",
                                  merge_ids=["REQ-015"], reviewer="QA")
    reqs = {r.requirement_id: r for r in repo.get_requirements(demo_run.run_id)}
    assert reqs["REQ-015"].review_status == ReviewStatus.REJECTED and "Merged into REQ-009" in reqs["REQ-015"].reviewer_note
    assert set(reqs["REQ-015"].source_statement_ids) <= set(keep.source_statement_ids)


def test_continue_run_uses_reviewed_requirements(demo_run, repo, session_id, settings):
    svc = ReviewService(repo)
    svc.apply(session_id=session_id, run_id=demo_run.run_id, item_type=ItemType.REQUIREMENT, item_id="REQ-006",
              action="reject")
    svc.apply(session_id=session_id, run_id=demo_run.run_id, item_type=ItemType.REQUIREMENT, item_id="REQ-009",
              action="approve", edits={"description": "The Employee receives an email when a claim is approved."})
    run = AnalysisOrchestrator(repo, settings, MockProvider()).continue_run(
        demo_run.run_id, find_issues=False, generate_stories=True, max_tests_per_story=1)
    assert run.status == RunStatus.COMPLETED and run.run_id == demo_run.run_id
    stories = repo.get_stories(run.run_id)
    assert not any("REQ-006" in s.requirement_ids for s in stories)
    assert any("email when a claim is approved" in s.capability for s in stories)
    assert all(sum(t.story_id == s.story_id for t in repo.get_scenarios(run.run_id)) <= 1 for s in stories)
    assert repo.get_findings(run.run_id)   # findings kept


# ------------------------------------------------------------------ traceability & comparison
def test_traceability_and_comparison(demo_run, repo, session_id):
    reqs = repo.get_requirements(demo_run.run_id)
    docs = {d.document_id: d for d in repo.list_documents(session_id)}
    rows = traceability(reqs, repo.get_stories(demo_run.run_id), repo.get_scenarios(demo_run.run_id),
                        repo.get_findings(demo_run.run_id), docs)
    by_id = {r.requirement_id: r for r in rows}
    sms = next(r for r in reqs if "SMS" in r.description)
    assert "No user story" in by_id[sms.requirement_id].gaps
    assert any("Unresolved contradiction" in r.gaps for r in rows)
    cmp = compare_meeting_vs_brd(reqs, repo.get_findings(demo_run.run_id), docs)
    cats = {c.category for c in cmp}
    assert {"Conflict", "Missing from BRD", "Not discussed"} <= cats
    qs = clarification_questions(repo.get_findings(demo_run.run_id), reqs,
                                 {s.statement_id: s for s in repo.get_statements(demo_run.run_id)})
    assert len({q.question for q in qs}) == len(qs) and not any(q.question.lower().startswith("demo") for q in qs)


# ------------------------------------------------------------------ exports
def _text(docx_bytes: bytes) -> str:
    """Document text in reading order (paragraphs and tables interleaved)."""
    from docx.table import Table
    d = Document(io.BytesIO(docx_bytes))
    parts = []
    for item in d.iter_inner_content():
        if isinstance(item, Table):
            for row in item.rows:
                parts += [c.text for c in row.cells]
        else:
            parts.append(item.text)
    return "\n".join(parts)


def test_brd_draft_and_final(demo_run, repo, session_id):
    svc = ReviewService(repo)
    svc.apply(session_id=session_id, run_id=demo_run.run_id, item_type=ItemType.REQUIREMENT, item_id="REQ-006",
              action="approve", reviewer="Nikhil")
    draft = _text(build_brd(load_bundle(repo, session_id, demo_run.run_id)))
    for heading in ("Document control", "1. Executive summary", "5. Requirements", "7. User stories",
                    "9. Open items", "10. Traceability matrix", "12. Source register", "13. Review and approval"):
        assert heading in draft
    assert "DEMO DATA" in draft and "REQ-001" in draft and "PENDING REVIEW" in draft
    final = _text(build_brd(load_bundle(repo, session_id, demo_run.run_id, final=True)))
    body = final.split("6. Assumptions")[0]
    assert "REQ-006" in body and "REQ-001" not in body.split("5. Requirements")[1]
    assert "Requirements not yet approved" in final
    assert "Nikhil" in final     # audit trail


def test_review_report_and_workbook_and_jira(demo_run, repo, session_id):
    b = load_bundle(repo, session_id, demo_run.run_id)
    report = _text(build_review_report(b))
    assert "Contradictions between sources" in report and "45 days" in report and "Questions to ask" in report
    wb = load_workbook(io.BytesIO(build_workbook(b)))
    assert wb.sheetnames[:3] == ["Summary", "Requirements", "Findings & Questions"]
    for s in ("Traceability", "Meeting vs BRD", "Statements", "Acceptance criteria"):
        assert s not in wb.sheetnames
    assert wb["Requirements"].max_row == len(b.requirements) + 1
    assert str(wb["Summary"]["B5"].value).startswith("=COUNTA(")
    rows = list(csv.DictReader(io.StringIO(build_jira_csv(b).decode("utf-8-sig"))))
    assert rows and rows[0]["Issue Type"] == "Story" and "Given" in rows[0]["Acceptance Criteria"]
    b.final = True
    assert list(csv.DictReader(io.StringIO(build_jira_csv(b).decode("utf-8-sig")))) == []  # nothing approved yet


# ------------------------------------------------------------------ evaluation
def test_evaluation_demo_recall(settings):
    from eval.run_eval import evaluate
    result = evaluate(MockProvider(), settings)
    assert result.recall == 1.0, result.markdown()
    assert all(ok for _, ok, _ in result.must_not), result.markdown()


def test_clarification_document_type_always_included(service, repo, session_id):
    from ria.pipeline.orchestrator import documents_for_mode
    _confirm_all(service, repo, session_id)
    r = service.ingest(session_id, "client_answers.txt", b"Q: Window? A: 45 days confirmed by Finance.",
                       DocumentType.CLARIFICATIONS)
    service.confirm_extraction(session_id, r.document.document_id)
    docs = repo.list_documents(session_id)
    for mode in InputMode:
        assert r.document.document_id in [d.document_id for d in documents_for_mode(mode, docs)]
