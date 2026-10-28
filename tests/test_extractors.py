from __future__ import annotations

import io

from docx import Document
from pypdf import PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from ria.extractors import get_extractor
from ria.models.enums import SegmentKind, WarningCode, WarningSeverity
from tests.conftest import SAMPLES


def _extract(ext: str, data: bytes, name: str | None = None):
    return get_extractor(ext).extract(data, name or f"file{ext}")


def test_txt_transcript_detected_with_timestamps_and_continuations():
    text = ("[00:00:05] Priya Shah: We need claim approval.\n"
            "Rahul: Finance approval above 25k.\n"
            "  and also an audit trail.\n"
            "00:01:10 Priya Shah: Agreed.\n")
    out = _extract(".txt", text.encode())
    assert out.detected_transcript
    assert [b.speaker for b in out.blocks] == ["Priya Shah", "Rahul", "Priya Shah"]
    assert out.blocks[0].timestamp_start == "00:00:05"
    assert out.blocks[1].text.endswith("and also an audit trail.")


def test_txt_document_labels_are_not_speakers():
    text = "Status: Draft\nOwner: BA team\nPriority: High\n\nThe system shall send emails."
    out = _extract(".txt", text.encode())
    assert not out.detected_transcript


def test_markdown_headings_lists_and_tables():
    md = "# Scope\nIntro text.\n\n- Item one\n- Item two\n\n## Rules\n| ID | Rule |\n|---|---|\n| R1 | Max 45 days |\n"
    out = _extract(".md", md.encode(), "doc.md")
    kinds = [b.kind for b in out.blocks]
    assert SegmentKind.LIST_ITEM in kinds and SegmentKind.TABLE_ROW in kinds
    row = next(b for b in out.blocks if b.kind == SegmentKind.TABLE_ROW)
    assert row.text == "ID: R1 | Rule: Max 45 days" and row.heading == "Rules"


def test_vtt_merges_same_speaker_cues_and_keeps_unknown_speaker():
    out = _extract(".vtt", (SAMPLES / "expense_portal_meeting.vtt").read_bytes())
    assert out.detected_transcript
    first = out.blocks[0]
    assert first.speaker == "Priya Shah" and "paper forms" in first.text  # two cues merged
    assert first.timestamp_start == "00:00:03" and first.timestamp_end == "00:00:14"
    assert any(b.speaker is None for b in out.blocks)


def test_srt_with_inline_speaker():
    srt = "1\n00:00:01,000 --> 00:00:04,000\nPriya: Hello team\n\n2\n00:00:04,500 --> 00:00:06,000\nRahul: Hi\n"
    out = _extract(".srt", srt.encode())
    assert [(b.speaker, b.text) for b in out.blocks] == [("Priya", "Hello team"), ("Rahul", "Hi")]


def test_docx_headings_lists_tables_in_order():
    out = _extract(".docx", (SAMPLES / "expense_portal_brd.docx").read_bytes())
    assert not out.detected_transcript
    rows = [b for b in out.blocks if b.kind == SegmentKind.TABLE_ROW]
    assert any("BR-01" in r.text and "60 days" in r.text for r in rows)
    assert all(r.heading for r in rows)
    assert any(b.kind == SegmentKind.LIST_ITEM for b in out.blocks)


def test_docx_teams_transcript_format():
    d = Document()
    for speaker, ts, text in [("Priya Shah", "0:03", "Start please."), ("Rahul Mehta", "0:10", "45 days."),
                              ("Priya Shah", "0:20", "Agreed."), ("Anita Desai", "0:31", "Email only.")]:
        d.add_paragraph(f"{speaker}   {ts}")
        d.add_paragraph(text)
    buf = io.BytesIO(); d.save(buf)
    out = _extract(".docx", buf.getvalue())
    assert out.detected_transcript
    assert [(b.speaker, b.timestamp_start) for b in out.blocks][:2] == [("Priya Shah", "0:03"), ("Rahul Mehta", "0:10")]


def test_corrupt_docx_returns_blocking_warning():
    out = _extract(".docx", b"not a zip file")
    assert out.warnings[0].code == WarningCode.CORRUPT_FILE
    assert out.warnings[0].severity == WarningSeverity.BLOCKING


def test_pdf_pages_headings_and_header_footer_removal():
    out = _extract(".pdf", (SAMPLES / "travel_expense_policy.pdf").read_bytes())
    texts = " ".join(b.text for b in out.blocks)
    assert "Northwind Travel and Expense Policy" not in texts and "Page 1 of 3" not in texts
    assert any(w.code == WarningCode.HEADERS_FOOTERS_REMOVED for w in out.warnings)
    deadline = next(b for b in out.blocks if "90 days" in b.text)
    assert deadline.page_number == 1 and deadline.heading == "2 Submission Deadline"


def test_image_only_pdf_is_blocking():
    buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=A4); c.showPage(); c.save()
    out = _extract(".pdf", buf.getvalue())
    assert out.warnings[0].code == WarningCode.IMAGE_ONLY_PDF
    assert not out.blocks


def test_encrypted_pdf_is_blocking():
    w = PdfWriter(); w.add_blank_page(200, 200); w.encrypt("secret")
    buf = io.BytesIO(); w.write(buf)
    out = _extract(".pdf", buf.getvalue())
    assert out.warnings[0].code == WarningCode.ENCRYPTED_PDF


def test_csv_rows_become_table_segments():
    out = _extract(".csv", (SAMPLES / "existing_user_stories.csv").read_bytes())
    assert len(out.blocks) == 3
    assert out.blocks[0].text.startswith("Story ID: US-101 | Title: Upload receipt")
    assert "Acceptance Criteria" not in out.blocks[2].text  # empty cell skipped


def test_xlsx_rows_with_sheet_heading():
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Backlog"
    ws.append(["ID", "Story"]); ws.append(["US-1", "As an Employee ..."]); ws.append([None, None])
    buf = io.BytesIO(); wb.save(buf)
    out = _extract(".xlsx", buf.getvalue())
    assert len(out.blocks) == 1 and out.blocks[0].heading == "Backlog"
