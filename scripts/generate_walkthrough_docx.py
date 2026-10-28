"""Generate a professional Word document (.docx) for the RIA End-to-End Walkthrough Guide."""
import os
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ART_DIR = Path(r"C:\Users\nvishwakarma\.gemini\antigravity\brain\11390f70-5297-4397-a0ec-850308b761c5")
WS_DIR = Path(r"c:\Users\nvishwakarma\Downloads\1\requirement-intelligence-assistant-co\docs")
IMG_DIR = ART_DIR / "screenshots"

TARGET_ART = ART_DIR / "RIA_End_to_End_Walkthrough_Guide.docx"
TARGET_WS = WS_DIR / "RIA_End_to_End_Walkthrough_Guide.docx"

PRIMARY_COLOR = RGBColor(0x1F, 0x4E, 0x79)    # Deep Navy
SECONDARY_COLOR = RGBColor(0x2E, 0x75, 0xB6)  # Blue Accent
DARK_TEXT = RGBColor(0x26, 0x26, 0x26)        # Dark Slate
MUTED_TEXT = RGBColor(0x59, 0x59, 0x59)       # Medium Gray
BG_LIGHT = "F2F4F7"


def shade_cell(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for m, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def create_callout_box(doc, text: str, title: str = "KEY TAKEAWAY"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Cm(16.5)
    shade_cell(cell, "EDF4FB")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Left border styling in XML
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    left_b = OxmlElement("w:left")
    left_b.set(qn("w:val"), "single")
    left_b.set(qn("w:sz"), "24")  # 3pt thick
    left_b.set(qn("w:space"), "0")
    left_b.set(qn("w:color"), "1F4E79")
    tc_borders.append(left_b)
    for b_name in ("top", "bottom", "right"):
        b = OxmlElement(f"w:{b_name}")
        b.set(qn("w:val"), "none")
        tc_borders.append(b)
    tc_pr.append(tc_borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = PRIMARY_COLOR

    r_text = p.add_run(text)
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = DARK_TEXT

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)


def add_image_with_caption(doc, img_name: str, caption: str, width_cm: float = 16.0):
    img_path = IMG_DIR / img_name
    if not img_path.exists():
        print(f"Warning: Image {img_path} not found!")
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run = p_img.add_run()
    run.add_picture(str(img_path), width=Cm(width_cm))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(14)
    r_cap = p_cap.add_run(f"Figure: {caption}")
    r_cap.italic = True
    r_cap.font.size = Pt(9)
    r_cap.font.color.rgb = MUTED_TEXT


def build_docx():
    doc = Document()

    # Configure Margins: 2.0 cm all around
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)  # A4

    # ------------------------------------------------------------- Header / Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Requirement Intelligence Assistant (RIA)")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(26)
    r_title.font.color.rgb = PRIMARY_COLOR

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Complete End-to-End User Flow & Operations Walkthrough Guide")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = SECONDARY_COLOR

    # Document Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    meta_widths = [Cm(4.5), Cm(12.0)]
    meta_rows = [
        ("Application Name:", "Requirement Intelligence Assistant (RIA)"),
        ("Target Roles:", "Business Analysts, Product Owners, QA Engineers, Enterprise Architects"),
        ("Documentation Scope:", "End-to-End Workflow: Ingestion → AI Analysis → Review → Stories & Tests → Export"),
        ("Walkthrough Video:", "ria_walkthrough_video.webm (Full resolution application recording included)")
    ]
    for idx, (label, val) in enumerate(meta_rows):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = meta_widths[0], meta_widths[1]
        shade_cell(c0, "F2F4F7")
        shade_cell(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = DARK_TEXT

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = DARK_TEXT

    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(8)
    p_sep.paragraph_format.space_after = Pt(8)

    # ------------------------------------------------------------- Executive Summary
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. Executive Summary & Solution Architecture")
    r.font.color.rgb = PRIMARY_COLOR
    r.font.size = Pt(16)
    r.bold = True

    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.space_after = Pt(8)
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.add_run(
        "The Requirement Intelligence Assistant (RIA) transforms raw and often conflicting business inputs—"
        "such as meeting audio/video transcripts, preliminary BRDs, Jira tickets, and regulatory policy PDFs—"
        "into structured, production-ready agile requirements. It automates cross-document contradiction detection, "
        "synthesizes atomic functional requirements, writes 4-section user stories with Gherkin acceptance criteria, "
        "and produces comprehensive test scenario matrices and multi-format exports."
    )

    create_callout_box(
        doc,
        "RIA runs entirely with local determinism, PII privacy scrubbing, and flexible LLM backends (xAI Grok, Ollama, "
        "OpenAI-compatible endpoints). Spoken meeting decisions are reconciled directly against written specifications to "
        "eliminate discrepancy risk before sprint kickoff.",
        "CORE VALUE PROPOSITION"
    )

    # ------------------------------------------------------------- Video Walkthrough Callout
    h2_vid = doc.add_heading(level=2)
    r = h2_vid.add_run("Video Walkthrough")
    r.font.color.rgb = SECONDARY_COLOR
    r.font.size = Pt(13)

    p_vid = doc.add_paragraph()
    p_vid.paragraph_format.space_after = Pt(6)
    p_vid.add_run(
        "A full live-action walkthrough recording has been generated to accompany this guide. "
        "It demonstrates every page transition, button click, review action, and export deliverable:"
    )

    p_vlink = doc.add_paragraph()
    p_vlink.paragraph_format.space_after = Pt(12)
    rv1 = p_vlink.add_run("🎬 Video File: ")
    rv1.bold = True
    rv2 = p_vlink.add_run("docs/videos/ria_walkthrough_video.webm\n")
    rv3 = p_vlink.add_run("(Direct path: c:\\Users\\nvishwakarma\\Downloads\\1\\requirement-intelligence-assistant-co\\docs\\videos\\ria_walkthrough_video.webm)")
    rv3.font.size = Pt(9)
    rv3.font.color.rgb = MUTED_TEXT

    # ------------------------------------------------------------- Step 1: Home Dashboard
    doc.add_page_break()
    h1_s1 = doc.add_heading(level=1)
    r = h1_s1.add_run("2. Home Dashboard & Session Workspace")
    r.font.color.rgb = PRIMARY_COLOR
    r.font.size = Pt(16)
    r.bold = True

    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_after = Pt(8)
    p_s1.paragraph_format.line_spacing = 1.15
    p_s1.add_run(
        "The Home Dashboard serves as the central control room for RIA. From this screen, team members manage "
        "working sessions, switch contexts between different product releases, and track historical analysis runs."
    )

    add_image_with_caption(
        doc,
        "01_home_dashboard.png",
        "Home Dashboard displaying Active Session status, Quick Navigation cards, and Session Management table."
    )

    p_s1_bullets = doc.add_paragraph()
    p_s1_bullets.paragraph_format.space_after = Pt(4)
    p_s1_bullets.add_run("Key Features & Capabilities:").bold = True

    features_s1 = [
        ("Active Session Banner: ", "Displays current session ID and timestamp with direct navigation cards into the workflow."),
        ("Session Switcher: ", "Quickly toggle between active sprint backlogs and client requirement packages."),
        ("Session History & Administration: ", "Create fresh sessions, duplicate previous baselines for delta comparisons, or purge completed runs with safety confirmations."),
        ("De-cluttered Interface: ", "Telemetry and backend AI configuration details are cleanly abstracted away to keep team focus on business analysis.")
    ]
    for b_title, b_desc in features_s1:
        p_b = doc.add_paragraph(style="List Bullet")
        p_b.paragraph_format.space_after = Pt(3)
        r_bt = p_b.add_run(b_title)
        r_bt.bold = True
        p_b.add_run(b_desc)

    # ------------------------------------------------------------- Step 2: Ingestion
    doc.add_page_break()
    h1_s2 = doc.add_heading(level=1)
    r = h1_s2.add_run("3. Document Ingestion (Prepare & Run)")
    r.font.color.rgb = PRIMARY_COLOR
    r.font.size = Pt(16)
    r.bold = True

    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_after = Pt(8)
    p_s2.paragraph_format.line_spacing = 1.15
    p_s2.add_run(
        "The Prepare & Run screen handles the ingestion and staging of raw project documentation. "
        "RIA supports heterogeneous inputs across text, markdown, Word documents, PDF manuals, and meeting subtitle tracks."
    )

    add_image_with_caption(
        doc,
        "02_prepare_and_run_documents.png",
        "Prepare & Run Document Ingestion showing multi-format document staging, categorization, and parsing status."
    )

    p_s2_bullets = doc.add_paragraph()
    p_s2_bullets.paragraph_format.space_after = Pt(4)
    p_s2_bullets.add_run("Ingestion Architecture:").bold = True

    features_s2 = [
        ("Multi-Format Parsing: ", "Accepts .txt, .md, .docx, .pdf, .vtt, .srt, and .csv files seamlessly."),
        ("Document Classification: ", "Categorize files as Meeting Transcript, BRD / PRD, Business Policy, or Architecture Spec."),
        ("Automated PII Redaction: ", "Sensitive names, emails, phone numbers, and financial details are anonymized before LLM transmission."),
        ("Document Inventory: ", "Clean table tracking Document ID (DOC-001), filenames, file sizes, and parsing status.")
    ]
    for b_title, b_desc in features_s2:
        p_b = doc.add_paragraph(style="List Bullet")
        p_b.paragraph_format.space_after = Pt(3)
        r_bt = p_b.add_run(b_title)
        r_bt.bold = True
        p_b.add_run(b_desc)

    # ------------------------------------------------------------- Step 3: AI Pipeline
    h1_s3 = doc.add_heading(level=1)
    r = h1_s3.add_run("4. AI Pipeline Execution (Prepare & Run)")
    r.font.color.rgb = PRIMARY_COLOR
    r.font.size = Pt(16)
    r.bold = True

    p_s3 = doc.add_paragraph()
    p_s3.paragraph_format.space_after = Pt(8)
    p_s3.paragraph_format.line_spacing = 1.15
    p_s3.add_run(
        "Once source documents are staged, the user configures the reconciliation mode and triggers the automated AI pipeline. "
        "The system processes the materials deterministically through five sequential analysis stages."
    )

    add_image_with_caption(
        doc,
        "03_prepare_and_run_analysis_options.png",
        "Analysis Configuration options and one-click execution trigger on Prepare & Run screen."
    )

    create_callout_box(
        doc,
        "Pipeline Stages: 1) Statements Extraction → 2) Requirements Deduction → 3) Cross-Document Conflict & Gap Detection → "
        "4) Agile User Story Formulation → 5) Acceptance Criteria & QA Test Derivation.",
        "PIPELINE AUTOMATION FLOW"
    )

    # ------------------------------------------------------------- Step 4: Review & Analysis
    doc.add_page_break()
    h1_s4 = doc.add_heading(level=1)
    r = h1_s4.add_run("5. Review & Analysis Workbench")
    r.font.color.rgb = PRIMARY_COLOR
    r.font.size = Pt(16)
    r.bold = True

    p_s4 = doc.add_paragraph()
    p_s4.paragraph_format.space_after = Pt(8)
    p_s4.paragraph_format.line_spacing = 1.15
    p_s4.add_run(
        "The Review & Analysis screen gives Business Analysts complete governance over generated requirements. "
        "Metric summary cards provide real-time status counts, led by the newly configured 'Confirm Requirement' card."
    )

    add_image_with_caption(
        doc,
        "04_review_and_analysis_overview.png",
        "Review & Analysis Workbench highlighting the 'Confirm Requirement' card and interactive requirements table."
    )

    features_s4 = [
        ("Confirm Requirement Card: ", "Instantly highlights requirements designated as confirmed and ready for sprint execution."),
        ("Status Lifecycle: ", "Track Pending Review, Approved (Ready for Stories), Needs Clarification, Conflict Detected, and Rejected."),
        ("Requirements Catalog: ", "Examines Requirement ID (REQ-001), category (Functional, Non-Functional, Security), title, and sources."),
        ("Inline Action Controls: ", "Approve, flag, request changes, or reject individual items with full audit trail logging.")
    ]
    for b_title, b_desc in features_s4:
        p_b = doc.add_paragraph(style="List Bullet")
        p_b.paragraph_format.space_after = Pt(3)
        r_bt = p_b.add_run(b_title)
        r_bt.bold = True
        p_b.add_run(b_desc)

    # ------------------------------------------------------------- Step 5: Clarification Email
    h1_s5 = doc.add_heading(level=1)
    r = h1_s5.add_run("6. Stakeholder Clarification Email")
    r.font.color.rgb = PRIMARY_COLOR
    r.font.size = Pt(16)
    r.bold = True

    p_s5 = doc.add_paragraph()
    p_s5.paragraph_format.space_after = Pt(8)
    p_s5.paragraph_format.line_spacing = 1.15
    p_s5.add_run(
        "When contradictions between spoken meetings and written BRDs are detected, RIA synthesizes an executive clarification "
        "email. All findings are consolidated into a single unified Conflicting Requirements table."
    )

    add_image_with_caption(
        doc,
        "05_review_and_analysis_email_draft.png",
        "Clarification Email Panel featuring consolidated Conflicting Requirements table and direct Outlook dispatch."
    )

    features_s5 = [
        ("Consolidated Conflicting Requirements Table: ", "Eliminates duplicate tables; combines conflict title, impact, meeting evidence, BRD text, and clarifying question."),
        ("Direct Outlook COM Integration: ", "Click '🚀 Open Directly in Outlook' to spawn a formatted draft email via Windows Outlook COM."),
        ("Auto-Subject & Recipients: ", "Pre-fills relevant project tags and formatted HTML tables ready for executive review.")
    ]
    for b_title, b_desc in features_s5:
        p_b = doc.add_paragraph(style="List Bullet")
        p_b.paragraph_format.space_after = Pt(3)
        r_bt = p_b.add_run(b_title)
        r_bt.bold = True
        p_b.add_run(b_desc)

    # ------------------------------------------------------------- Step 6: Stories & Tests
    doc.add_page_break()
    h1_s6 = doc.add_heading(level=1)
    r = h1_s6.add_run("7. Agile Stories & QA Test Scenarios")
    r.font.color.rgb = PRIMARY_COLOR
    r.font.size = Pt(16)
    r.bold = True

    p_s6 = doc.add_paragraph()
    p_s6.paragraph_format.space_after = Pt(8)
    p_s6.paragraph_format.line_spacing = 1.15
    p_s6.add_run(
        "The Stories & Tests module translates approved requirements into structured engineering backlogs. "
        "Every story follows the strict 4-section architecture specified by development leads."
    )

    add_image_with_caption(
        doc,
        "06_stories_and_tests_board.png",
        "Stories & Tests Board illustrating the 4-section layout with complete Functional Requirement text."
    )

    p_s6_arch = doc.add_paragraph()
    p_s6_arch.paragraph_format.space_after = Pt(4)
    p_s6_arch.add_run("The 4-Section User Story Architecture:").bold = True

    sections_s6 = [
        ("1. Description: ", "Standard agile narrative ('As a [User], I want to [Action], so that [Benefit]')."),
        ("2. Path or Navigation: ", "Explicit system navigation route (e.g. 'Customer Portal → Claims → Settlement Details')."),
        ("3. Functional Requirement: ", "The unabridged, complete text of the underlying requirement specifications."),
        ("4. Acceptance Requirement & QA Tests: ", "Gherkin Given-When-Then criteria with inline edit popovers, paired with positive Happy Path and negative Edge Case QA test scenarios.")
    ]
    for b_title, b_desc in sections_s6:
        p_b = doc.add_paragraph(style="List Bullet")
        p_b.paragraph_format.space_after = Pt(3)
        r_bt = p_b.add_run(b_title)
        r_bt.bold = True
        p_b.add_run(b_desc)

    # ------------------------------------------------------------- Step 7: Jira Push
    h1_s7 = doc.add_heading(level=1)
    r = h1_s7.add_run("8. Direct Jira Integration")
    r.font.color.rgb = PRIMARY_COLOR
    r.font.size = Pt(16)
    r.bold = True

    p_s7 = doc.add_paragraph()
    p_s7.paragraph_format.space_after = Pt(8)
    p_s7.paragraph_format.line_spacing = 1.15
    p_s7.add_run(
        "The built-in Jira Workbench allows engineering teams to push sprint-ready stories and QA scenarios directly into "
        "Jira Cloud or Server with zero manual re-entry."
    )

    add_image_with_caption(
        doc,
        "07_stories_and_tests_jira.png",
        "Jira Push tab showing live API connection test, project configuration, and export settings."
    )

    features_s7 = [
        ("Direct REST API Sync: ", "Connects directly via secure API tokens to Jira Cloud / Data Center."),
        ("Connection Test Tool: ", "Instant feedback verifying Jira project key accessibility and issue creation permissions."),
        ("Issue Mapping: ", "Maps descriptions, acceptance criteria, test cases, and priority tags into Jira fields."),
        ("Hyperlinked Backlog: ", "Captures created Jira keys (e.g. PROJ-102) with direct clickable links into Jira boards.")
    ]
    for b_title, b_desc in features_s7:
        p_b = doc.add_paragraph(style="List Bullet")
        p_b.paragraph_format.space_after = Pt(3)
        r_bt = p_b.add_run(b_title)
        r_bt.bold = True
        p_b.add_run(b_desc)

    # ------------------------------------------------------------- Step 8: Export Center
    doc.add_page_break()
    h1_s8 = doc.add_heading(level=1)
    r = h1_s8.add_run("9. Multi-Format Export Center")
    r.font.color.rgb = PRIMARY_COLOR
    r.font.size = Pt(16)
    r.bold = True

    p_s8 = doc.add_paragraph()
    p_s8.paragraph_format.space_after = Pt(8)
    p_s8.paragraph_format.line_spacing = 1.15
    p_s8.add_run(
        "The Export Center packages all validated requirements, stories, tests, and conflict logs into official project artifacts. "
        "The top summary displays requirement and user story counts, followed by download buttons for each format."
    )

    add_image_with_caption(
        doc,
        "08_export_center.png",
        "Export Center displaying Requirement & User Story metric cards and multi-format download options."
    )

    # Table of Deliverables
    deliv_table = doc.add_table(rows=4, cols=3)
    deliv_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    deliv_table.autofit = False
    deliv_widths = [Cm(3.5), Cm(4.0), Cm(9.0)]

    headers = ["Format", "File Type", "Content & Structure"]
    for i, h in enumerate(headers):
        c = deliv_table.rows[0].cells[i]
        c.width = deliv_widths[i]
        shade_cell(c, "1F4E79")
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    deliv_data = [
        ("Excel Workbook", ".xlsx", "Summary sheet with dynamic count formulas, Requirements sheet, unified Findings & Questions sheet, User Stories sheet, and Test Scenarios sheet."),
        ("Word BRD", ".docx", "Formatted Business Requirements Document with executive summary, atomic requirement definitions, conflict resolution log, and stakeholder sign-off section."),
        ("Jira CSV", ".csv", "Pre-formatted standard CSV ready for bulk backlog import into Jira Cloud, Server, or Azure DevOps.")
    ]

    for idx, (fmt, ftype, desc) in enumerate(deliv_data, start=1):
        row = deliv_table.rows[idx]
        for col_idx, text_val in enumerate((fmt, ftype, desc)):
            cell = row.cells[col_idx]
            cell.width = deliv_widths[col_idx]
            shade_cell(cell, "FFFFFF" if idx % 2 else "F9FAFB")
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(text_val)
            r.font.size = Pt(9.0)
            r.font.color.rgb = DARK_TEXT

    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(14)
    p_end.add_run("Summary Checklist for Product Teams:").bold = True

    # Summary table
    sum_table = doc.add_table(rows=6, cols=3)
    sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sum_table.autofit = False
    sum_widths = [Cm(3.0), Cm(6.5), Cm(7.0)]

    s_headers = ["Phase", "Activity", "RIA Screen"]
    for i, h in enumerate(s_headers):
        c = sum_table.rows[0].cells[i]
        c.width = sum_widths[i]
        shade_cell(c, "2E75B6")
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    sum_data = [
        ("1. Ingestion", "Upload transcripts, PRDs, and policy docs", "Prepare & Run"),
        ("2. Analysis", "Run cross-document AI reconciliation", "Prepare & Run"),
        ("3. Governance", "Validate requirements & verify Confirm Requirement", "Review & Analysis"),
        ("4. Clarification", "Review merged conflicts & open draft in Outlook", "Review & Analysis"),
        ("5. Backlog & Export", "Inspect 4-section stories, push Jira & export", "Stories & Tests / Export")
    ]
    for idx, (ph, act, scr) in enumerate(sum_data, start=1):
        row = sum_table.rows[idx]
        for col_idx, text_val in enumerate((ph, act, scr)):
            cell = row.cells[col_idx]
            cell.width = sum_widths[col_idx]
            shade_cell(cell, "FFFFFF" if idx % 2 else "F9FAFB")
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(text_val)
            r.font.size = Pt(9.0)
            r.font.color.rgb = DARK_TEXT

    # Save to artifact directory and workspace docs
    doc.save(str(TARGET_ART))
    doc.save(str(TARGET_WS))
    print(f"Document successfully saved to:\n- {TARGET_ART}\n- {TARGET_WS}")


if __name__ == "__main__":
    build_docx()
