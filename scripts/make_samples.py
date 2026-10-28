"""Generate the synthetic sample documents in sample_data/.

All content is fictional ("Northwind" is a made-up company). The documents contain
deliberately planted issues listed in sample_data/expected_findings.json, used by
the evaluation script in Phase 5.

Run:  python scripts/make_samples.py
"""
from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent.parent / "sample_data"

VTT_CUES = [
    ("00:00:03.120", "00:00:09.400", "Priya Shah", "Okay everyone, let's start. Today we are finalising the requirements for the new expense reimbursement portal."),
    ("00:00:09.400", "00:00:14.800", "Priya Shah", "The goal is to cut reimbursement time and stop the paper forms."),
    ("00:00:15.200", "00:00:24.900", "Rahul Mehta", "From finance side, employees must submit an expense clam within 30 days of the expense date. That is decided."),
    ("00:00:25.300", "00:00:33.000", "Anita Desai", "Haan, theek hai, but manager approval zaroori hai for anything above 500."),
    ("00:00:33.400", "00:00:41.700", "Priya Shah", "Agreed. And approvals should happen quickly, people complain about delays."),
    ("00:00:42.000", "00:00:50.500", "Rahul Mehta", "Claims above twenty five thousand also need finance approval after the manager."),
    ("00:00:51.000", "00:00:58.200", "Anita Desai", "Maybe we could also send an SMS when the claim is approved."),
    ("00:00:58.600", "00:01:06.300", "Priya Shah", "Email notification is definitely required. SMS we can discuss later."),
    ("00:01:06.800", "00:01:15.000", None, "What happens if the manager rejects it, can the employee resubmit?"),
    ("00:01:15.400", "00:01:21.900", "Priya Shah", "Good question. We'll figure that out later."),
    ("00:01:22.300", "00:01:31.800", "Rahul Mehta", "Sorry, correction. I said 30 days earlier, it should be 45 days. Finance confirmed this morning."),
    ("00:01:32.200", "00:01:41.500", "Anita Desai", "Receipts must be uploaded as PDF or image. And the GL code has to be selected by the employee."),
    ("00:01:42.000", "00:01:50.100", "Rahul Mehta", "I will share the cost centre list by Friday. Send questions to rahul.mehta@northwind.example."),
    ("00:01:50.500", "00:01:57.000", "Priya Shah", "Mileage claims are out of scope for phase one. Thanks everyone."),
]


def make_vtt() -> None:
    lines = ["WEBVTT", ""]
    for i, (start, end, speaker, text) in enumerate(VTT_CUES, start=1):
        lines.append(f"{i}-0")
        lines.append(f"{start} --> {end}")
        lines.append(f"<v {speaker}>{text}</v>" if speaker else text)
        lines.append("")
    (OUT / "expense_portal_meeting.vtt").write_text("\n".join(lines), encoding="utf-8")


def make_brd() -> None:
    d = Document()
    d.add_heading("Business Requirements Document - Expense Reimbursement Portal", 0)
    d.add_paragraph("Version 0.3 (Draft). Owner: Business Analysis Team, Northwind Ltd. (fictional).")

    d.add_heading("1. Purpose", 1)
    d.add_paragraph("Northwind employees currently submit paper expense forms. Reimbursement takes on "
                    "average 21 days. The new portal will digitise submission, approval and payment tracking.")

    d.add_heading("2. Scope", 1)
    d.add_paragraph("Online submission of travel and meal expense claims.", style="List Bullet")
    d.add_paragraph("Manager and finance approval workflow.", style="List Bullet")
    d.add_paragraph("Email notifications for claim status changes.", style="List Bullet")
    d.add_paragraph("Out of scope: mileage claims and corporate card reconciliation.")

    d.add_heading("3. Stakeholders", 1)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Role", "Responsibility"
    for role, resp in [("Employee", "Submits expense claims"),
                       ("Line Manager", "Approves or rejects claims of direct reports"),
                       ("Finance Approver", "Approves high-value claims and releases payment")]:
        r = t.add_row().cells
        r[0].text, r[1].text = role, resp

    d.add_heading("4. Business Requirements", 1)
    t = d.add_table(rows=1, cols=3)
    for c, h in zip(t.rows[0].cells, ("ID", "Requirement", "Priority")):
        c.text = h
    for row in [
        ("BR-01", "Employees must submit a claim within 60 days of the expense date.", "High"),
        ("BR-02", "A receipt must be attached for every expense line above INR 500.", "High"),
        ("BR-03", "Claims above INR 1,000 require Line Manager approval.", "High"),
        ("BR-04", "Claims above INR 25,000 require Finance Approver approval after the Line Manager.", "High"),
        ("BR-05", "The employee receives an email when a claim is approved or rejected.", "Medium"),
        ("BR-06", "The employee must select a GL code for every expense line.", "Medium"),
        ("BR-07", "The system should be user-friendly.", "Low"),
    ]:
        r = t.add_row().cells
        for c, v in zip(r, row):
            c.text = v

    d.add_heading("5. Non-Functional Requirements", 1)
    d.add_paragraph("The portal should load fast and be available during business hours.")
    d.add_paragraph("All approval actions must be recorded in an audit trail.")

    d.add_heading("6. Open Questions", 1)
    d.add_paragraph("Payment integration with the payroll system is to be confirmed.")
    d.save(OUT / "expense_portal_brd.docx")


def make_policy_pdf() -> None:
    pages = [
        ["1 Applicability",
         "This policy applies to all permanent employees travelling on company business.",
         "2 Submission Deadline",
         "Expense claims older than 90 days from the expense date will not be reimbursed.",
         "3 Receipts",
         "Original receipts are required for any single expense above INR 500.",
         "Digital copies are accepted when the original is lost, with manager confirmation."],
        ["4 Meal Allowance",
         "Meal expenses are reimbursed up to INR 1,500 per day in metro cities.",
         "Meal expenses are reimbursed up to INR 1,000 per day in other cities.",
         "5 Approvals",
         "All claims require approval from the employee's line manager before payment.",
         "Claims above INR 25,000 additionally require approval from Finance."],
        ["6 Payment",
         "Approved claims are paid with the next monthly payroll cycle.",
         "7 Exceptions",
         "Exceptions to this policy require written approval from the Head of Finance."],
    ]
    c = canvas.Canvas(str(OUT / "travel_expense_policy.pdf"), pagesize=A4)
    w, h = A4
    for i, lines in enumerate(pages, start=1):
        c.setFont("Helvetica", 9)
        c.drawString(50, h - 40, "Northwind Travel and Expense Policy v3.2 - Internal")
        c.drawString(w - 110, 30, f"Page {i} of {len(pages)}")
        y = h - 90
        for ln in lines:
            heading = ln[0].isdigit() and len(ln) < 40
            c.setFont("Helvetica-Bold" if heading else "Helvetica", 12 if heading else 11)
            if heading:
                y -= 8
            c.drawString(50, y, ln)
            y -= 22
        c.showPage()
    c.save()


def make_csv() -> None:
    rows = [
        ("Story ID", "Title", "Story", "Acceptance Criteria", "Status"),
        ("US-101", "Upload receipt", "As an Employee, I want to upload a receipt image so that my claim has proof of expense.",
         "Given a claim line, when I upload a JPG or PDF, then it is attached to the line.", "Ready"),
        ("US-102", "Manager approval", "As a Line Manager, I want to approve or reject claims from my team so that valid expenses are reimbursed.",
         "Given a submitted claim, when I approve it, then the status changes to Approved.", "Ready"),
        ("US-103", "Claim history", "As an Employee, I want to see my claim history so that I can track payments.",
         "", "Draft"),
    ]
    with open(OUT / "existing_user_stories.csv", "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    make_vtt(); make_brd(); make_policy_pdf(); make_csv()
    print("Sample documents written to", OUT)
