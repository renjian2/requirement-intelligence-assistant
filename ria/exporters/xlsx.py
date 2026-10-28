"""Excel workbook export: one sheet per artefact, filterable, with a formula-based summary."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ria.exporters.bundle import ExportBundle
from ria.models.enums import FINDING_LABELS
from ria.services.traceability import clarification_questions

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
THIN = Side(style="thin", color="BFBFBF")
STATUS_FILL = {"APPROVED": "C6EFCE", "EDITED": "C6EFCE", "RESOLVED": "C6EFCE", "REJECTED": "F2F2F2",
               "OUT_OF_SCOPE": "F2F2F2", "NEEDS_CLARIFICATION": "FFEB9C", "PENDING": "FFFFFF",
               "HIGH": "F8CBAD", "MEDIUM": "FFE699", "LOW": "DDEBF7", "CONFLICTING": "F8CBAD",
               "UNSUPPORTED": "D9D9D9", "AMBIGUOUS": "FFE699", "PROPOSED": "FFF2CC"}


def _j(v) -> str:
    if isinstance(v, list):
        return "\n".join(str(x) for x in v)
    return "" if v is None else str(getattr(v, "value", v))


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int],
           color_cols: tuple[str, ...] = ()) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([_j(v) for v in row])
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    color_idx = [headers.index(c) + 1 for c in color_cols if c in headers]
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        for i in color_idx:
            fill = STATUS_FILL.get(str(r[i - 1].value))
            if fill:
                r[i - 1].fill = PatternFill("solid", fgColor=fill)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def build_workbook(b: ExportBundle) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    docname = {d: x.file_name for d, x in b.docs.items()}

    reqs = b.requirements
    _sheet(wb, "Requirements",
           ["ID", "Area", "Title", "Description", "Type", "Actors", "Trigger", "Expected behaviour", "Business rules",
            "Exceptions", "Priority", "Priority source", "Evidence", "Confidence", "Review", "Reviewer note",
            "Sources", "Segments", "Statements", "Open questions", "Assumptions", "Validation notes"],
           [[r.requirement_id, r.group, r.title, r.description, r.requirement_type, r.actors, r.trigger,
             r.expected_behavior, r.business_rules, r.exceptions, r.priority, r.priority_source, r.evidence_status,
             r.confidence, r.review_status, r.reviewer_note, [docname.get(d, d) for d in r.document_ids],
             ", ".join(r.segment_ids), ", ".join(r.source_statement_ids), r.unresolved_questions, r.assumptions,
             r.validation_notes] for r in reqs],
           [10, 18, 30, 50, 14, 18, 22, 30, 30, 24, 10, 13, 14, 11, 14, 24, 28, 24, 20, 40, 30, 40],
           ("Evidence", "Review"))
    finding_ids = {f.finding_id for f in b.findings}
    questions = clarification_questions(b.findings, reqs, b.statements)
    fq_rows = []
    for f in b.findings:
        fq_rows.append([
            f.finding_id, FINDING_LABELS[f.finding_type], f.severity, f.title, f.description,
            f.clarification_question, f.ask_whom, ", ".join(f.requirement_ids),
            [f"{v.requirement_id}: {v.value}" for v in f.values],
            f.business_impact, f.dev_impact, f.qa_impact, f.suggested_resolution,
            f.detected_by, f.review_status, f.resolution, ", ".join(f.segment_ids),
        ])
    extra_q_idx = 1
    for q in questions:
        if q.source not in finding_ids:
            fq_rows.append([
                f"Q-{extra_q_idx:03d}", "Open Question", q.severity, f"Question for {q.ask_whom}",
                q.question, q.question, q.ask_whom, ", ".join(q.requirement_ids),
                [], "", "", "", "", "CODE", "PENDING", "", "",
            ])
            extra_q_idx += 1
    _sheet(wb, "Findings & Questions",
           ["ID", "Type", "Severity", "Title", "Description", "Clarification question", "Ask whom",
            "Requirements", "Values", "Business impact", "Dev impact", "QA impact", "Suggested resolution",
            "Detected by", "Review", "Resolution", "Evidence segments"],
           fq_rows,
           [10, 18, 10, 34, 50, 50, 22, 18, 30, 26, 26, 26, 34, 11, 14, 30, 30],
           ("Severity", "Review"))
    _sheet(wb, "User stories",
           ["ID", "Area", "Title", "Story", "Actor", "Requirements", "Priority", "Acceptance criteria", "Review",
            "Notes"],
           [[s.story_id, s.group, s.title, s.story_text, s.actor, ", ".join(s.requirement_ids), s.priority,
             [f"{a.ac_id}: {a.text}" for a in s.acceptance_criteria], s.review_status,
             s.validation_notes + s.clarification_questions] for s in b.stories],
           [9, 18, 30, 60, 18, 16, 10, 80, 14, 40], ("Review",))
    _sheet(wb, "Test scenarios",
           ["ID", "Story", "Requirements", "Title", "Type", "Priority", "Preconditions", "Steps", "Expected result",
            "Review", "Notes"],
           [[t.scenario_id, t.story_id, ", ".join(t.requirement_ids), t.title, t.scenario_type, t.priority,
             t.preconditions, [f"{i}. {x}" for i, x in enumerate(t.steps, start=1)], t.expected_result,
             t.review_status, t.validation_notes] for t in b.scenarios],
           [9, 9, 16, 40, 16, 10, 34, 50, 44, 14, 30], ("Priority", "Review"))
    _sheet(wb, "Sources", ["Document", "File", "Type", "Segments", "Words", "Status", "SHA-256 (start)"],
           [[d.document_id, d.file_name, d.document_type, d.stats.segments, d.stats.words, d.status, d.sha256[:16]]
            for d in b.docs.values()], [10, 36, 22, 10, 10, 12, 20])

    # Summary with live formulas (counts update if the reviewer edits statuses in Excel).
    ws = summary
    ws["A1"] = f"Requirement Intelligence Assistant – {b.session_name}"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = (f"Session: {b.session_name} · {b.run.created_at:%Y-%m-%d %H:%M} UTC"
                + (" · DEMO DATA" if b.run.is_demo else ""))
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")
    lines = [
        ("Requirements", "=COUNTA(Requirements!A:A)-1"),
        ("  Approved", '=COUNTIF(Requirements!O:O,"APPROVED")+COUNTIF(Requirements!O:O,"EDITED")'),
        ("  Pending review", '=COUNTIF(Requirements!O:O,"PENDING")'),
        ("  Needs clarification", '=COUNTIF(Requirements!O:O,"NEEDS_CLARIFICATION")'),
        ("  Rejected", '=COUNTIF(Requirements!O:O,"REJECTED")+COUNTIF(Requirements!O:O,"OUT_OF_SCOPE")'),
        ("User stories", "=COUNTA('User stories'!A:A)-1"),
        ("Test scenarios", "=COUNTA('Test scenarios'!A:A)-1"),
    ]
    ws["A4"], ws["B4"] = "Metric", "Count"
    for c in ("A4", "B4"):
        ws[c].font = Font(name=FONT, bold=True, color="FFFFFF")
        ws[c].fill = HEADER_FILL
    for i, (label, formula) in enumerate(lines, start=5):
        c1 = ws.cell(i, 1, label)
        c2 = ws.cell(i, 2, formula)
        c1.font = Font(name=FONT, bold=(not label.startswith("  ")))
        c2.font = Font(name=FONT)
        c1.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        c2.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 12
    note = 5 + len(lines) + 1
    ws.cell(note, 1, "Counts are live formulas over the other sheets. Evidence and confidence were checked by code; "
                     "review statuses come from human reviewers.").font = Font(name=FONT, size=9, italic=True)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 12
    note = 5 + len(lines) + 1
    ws.cell(note, 1, "Counts are live formulas over the other sheets. Evidence and confidence were checked by code; "
                     "review statuses come from human reviewers.").font = Font(name=FONT, size=9, italic=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
