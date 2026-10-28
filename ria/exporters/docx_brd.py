"""Business Requirements Document (DOCX) and Review Report, assembled by code from validated data.

Nothing in these documents is free text written by the AI at export time: every requirement,
story and finding comes from the stored, validated records, and missing sections say so.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import datetime, timezone

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ria.exporters.bundle import ExportBundle
from ria.models.enums import (
    DOCUMENT_TYPE_LABELS, EXCLUDED_REVIEW, FINDING_LABELS, INPUT_MODE_LABELS, REQUIREMENT_DOC_TYPES, SEVERITY_ORDER,
    FindingType, RequirementType, ReviewStatus, StatementType,
)
from ria.services.traceability import clarification_questions, compare_meeting_vs_brd, traceability

NOT_FOUND = "Not identified in the supplied source documents."
BLUE = RGBColor(0x1F, 0x4E, 0x79)
RED = RGBColor(0xC0, 0x00, 0x00)
TYPE_TITLES = [
    (RequirementType.BUSINESS, "Business requirements and rules"), (RequirementType.FUNCTIONAL, "Functional requirements"),
    (RequirementType.NON_FUNCTIONAL, "Non-functional requirements"), (RequirementType.DATA, "Data requirements"),
    (RequirementType.REPORTING, "Reporting requirements"), (RequirementType.INTEGRATION, "Integration requirements"),
    (RequirementType.SECURITY, "Security requirements"), (RequirementType.COMPLIANCE, "Compliance requirements"),
    (RequirementType.ACCESSIBILITY, "Accessibility requirements"),
]


# ---------------------------------------------------------------- low-level helpers
def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()  # noqa: SLF001
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()  # noqa: SLF001
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        elif text == instr:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = instr
        else:
            el = OxmlElement("w:t"); el.text = text
        run._r.append(el)  # noqa: SLF001


def _table(doc, headers: list[str], rows: list[list[str]], widths_cm: list[float], font_size: int = 8) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold, run.font.size = True, Pt(font_size)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, "1F4E79")
    _repeat_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            lines = str(v if v is not None else "").split("\n")
            p = cells[i].paragraphs[0]
            for k, line in enumerate(lines):
                if k:
                    p = cells[i].add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.add_run(line).font.size = Pt(font_size)
    # Fixed layout + grid widths + cell widths: needed for Word AND LibreOffice to respect widths.
    tbl_pr = t._tbl.tblPr  # noqa: SLF001
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = t._tbl.tblGrid  # noqa: SLF001
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths_cm):
            gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))
    for i, w in enumerate(widths_cm):
        t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    doc.add_paragraph()


def _para(doc, text: str, *, italic=False, bold=False, color=None, size=None, style=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.italic, run.bold = italic, bold
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return p


def _bullets(doc, items: list[str]) -> None:
    if not items:
        _para(doc, NOT_FOUND, italic=True)
        return
    for x in items:
        doc.add_paragraph(x, style="List Bullet")


def _setup(doc, title: str, draft_banner: str | None, landscape: bool = False) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        st = styles[name]
        st.font.name, st.font.size, st.font.color.rgb = "Arial", Pt(size), BLUE
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    hp = sec.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run(f"{title}" + (f"   ·   {draft_banner}" if draft_banner else ""))
    r.font.size = Pt(8)
    r.font.color.rgb = RED if draft_banner else RGBColor(0x59, 0x59, 0x59)
    r.bold = bool(draft_banner)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ").font.size = Pt(8)
    _field(fp, "PAGE")
    fp.add_run(" of ").font.size = Pt(8)
    _field(fp, "NUMPAGES")


def _status_label(item) -> str:
    return {"APPROVED": "Approved", "EDITED": "Approved (edited)", "PENDING": "PENDING REVIEW",
            "NEEDS_CLARIFICATION": "NEEDS CLARIFICATION"}.get(item.review_status.value, item.review_status.value)


def _sources(b: ExportBundle, segment_ids: list[str], limit: int = 3) -> str:
    out = []
    for sid in segment_ids[:limit]:
        seg = b.segments.get(sid)
        if seg:
            doc = b.docs.get(seg.document_id)
            loc = seg.location
            out.append(f"{doc.file_name if doc else seg.document_id}{' · ' + loc if loc else ''}")
    more = len(segment_ids) - limit
    return "\n".join(out) + (f"\n+{more} more" if more > 0 else "")


def _title_desc(title: str, desc: str) -> str:
    """Avoid repeating the description when the title is just its first words."""
    t = title.rstrip("…. ").lower()
    return desc if desc.lower().startswith(t) else f"{title}\n{desc}"


def _primary_understanding(b: ExportBundle):
    us = b.understandings
    for d, doc in b.docs.items():
        if doc.document_type in REQUIREMENT_DOC_TYPES and d in us:
            return us[d]
    return next(iter(us.values()), None)


# ---------------------------------------------------------------- BRD
def build_brd(b: ExportBundle) -> bytes:
    doc = Document()
    pending = b.pending_count
    banner = None
    if b.run.is_demo:
        banner = "DEMO DATA – not produced by AI"
    elif not b.final or pending:
        banner = "DRAFT – contains unapproved items" if pending else "DRAFT"
    _setup(doc, f"BRD – {b.session_name}", banner)
    reqs = b.included(b.requirements)
    stories = b.included(b.stories)
    scenarios = b.included(b.scenarios)
    now = datetime.now(timezone.utc)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run("Business Requirements Document")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(24), BLUE
    _para(doc, b.session_name, size=15, bold=True)
    if banner:
        _para(doc, banner + ". Items marked PENDING REVIEW have not been approved by a person.", bold=True, color=RED)

    doc.add_heading("Document control", level=1)
    _table(doc, ["Field", "Value"], [
        ["Status", "Final (approved items only)" if b.final else "Draft (all items not rejected)"],
        ["Generated", f"{now:%d %b %Y %H:%M} UTC by Requirement Intelligence Assistant"],
        ["Analysis run", f"{b.run.run_id} · {b.run.provider}/{b.run.model or '-'} · "
                         f"{INPUT_MODE_LABELS[b.run.input_mode]}"],
        ["Source documents", "\n".join(f"{d.document_id} {d.file_name}" for d in b.docs.values())],
        ["Review status", f"{sum(1 for x in b.requirements if x.review_status.value in ('APPROVED', 'EDITED'))} of "
                          f"{len(b.requirements)} requirements approved · {pending} item(s) pending"],
    ], [4.5, 12.5], 9)

    u = _primary_understanding(b)
    doc.add_heading("1. Executive summary", level=1)
    if u and (u.problem_statement or u.objectives):
        if u.problem_statement:
            _para(doc, u.problem_statement)
        if u.domain:
            _para(doc, f"Domain: {u.domain}.")
    else:
        _para(doc, NOT_FOUND, italic=True)
    open_f = [f for f in b.findings if f.review_status not in (ReviewStatus.REJECTED, ReviewStatus.RESOLVED)]
    high = sum(1 for f in open_f if f.severity.value == "HIGH")
    _para(doc, f"This document contains {len(reqs)} requirement(s), {len(stories)} user stor"
               f"{'y' if len(stories) == 1 else 'ies'} and {len(scenarios)} important test scenario(s). "
               f"{len(open_f)} open issue(s) remain, {high} of them high severity – see section 9.")

    doc.add_heading("2. Business problem and objectives", level=1)
    objectives = [o for x in b.understandings.values() for o in x.objectives]
    objectives += [s.text for s in b.statements.values() if s.statement_type == StatementType.OBJECTIVE]
    _bullets(doc, list(dict.fromkeys(objectives)))

    doc.add_heading("3. Scope", level=1)
    doc.add_heading("In scope (functional areas)", level=2)
    _bullets(doc, list(dict.fromkeys(r.group.split(" (part")[0] for r in reqs)))
    doc.add_heading("Out of scope", level=2)
    out = [s.text for s in b.statements.values() if s.statement_type == StatementType.OUT_OF_SCOPE]
    out += [f"{r.requirement_id}: {r.title} (marked out of scope by reviewer)" for r in b.requirements
            if r.review_status == ReviewStatus.OUT_OF_SCOPE]
    _bullets(doc, out)

    doc.add_heading("4. Stakeholders and actors", level=1)
    actors = list(dict.fromkeys([a for x in b.understandings.values() for a in x.actors] + [a for r in reqs for a in r.actors]))
    stakeholders = list(dict.fromkeys(a for x in b.understandings.values() for a in x.stakeholders))
    speakers = sorted({s.speaker for s in b.segments.values() if s.speaker})
    rows = [["Actors (use the system)", ", ".join(actors) or NOT_FOUND],
            ["Stakeholders", ", ".join(stakeholders) or NOT_FOUND],
            ["Meeting participants", ", ".join(speakers) or NOT_FOUND]]
    _table(doc, ["Group", "Names / roles"], rows, [5, 12], 9)

    doc.add_heading("5. Requirements", level=1)
    sec = 0
    for rtype, title in TYPE_TITLES:
        items = [r for r in reqs if r.requirement_type == rtype]
        if not items:
            continue
        sec += 1
        doc.add_heading(f"5.{sec} {title}", level=2)
        _table(doc, ["ID", "Requirement", "Actors", "Priority", "Evidence", "Review", "Source"],
               [[r.requirement_id, _title_desc(r.title, r.description) + (
                   f"\nRules: {'; '.join(r.business_rules)}" if r.business_rules else "") + (
                   f"\nExceptions: {'; '.join(r.exceptions)}" if r.exceptions else ""),
                 ", ".join(r.actors) or "-",
                 (f"{r.priority.value}" + (" (stated)" if r.priority_source.value == "STATED" else " (suggested)"))
                 if r.priority else "-",
                 f"{r.evidence_status.value}\n{r.confidence}/100", _status_label(r), _sources(b, r.segment_ids)]
                for r in items], [1.6, 7.0, 2.2, 1.6, 1.8, 1.8, 3.0])
    if not sec:
        _para(doc, NOT_FOUND, italic=True)

    doc.add_heading("6. Assumptions, constraints, dependencies and risks", level=1)
    for label, types, extra in (
            ("Assumptions", {StatementType.ASSUMPTION}, [a for r in reqs for a in r.assumptions]),
            ("Constraints", {StatementType.CONSTRAINT}, []),
            ("Dependencies", {StatementType.DEPENDENCY}, []),
            ("Risks", {StatementType.RISK}, [])):
        doc.add_heading(label, level=2)
        _bullets(doc, list(dict.fromkeys([s.text for s in b.statements.values() if s.statement_type in types] + extra)))

    doc.add_heading("7. User stories and acceptance criteria", level=1)
    if not stories:
        _para(doc, "No user stories generated yet.", italic=True)
    for s in stories:
        doc.add_heading(f"{s.story_id} · {s.title}", level=3)
        _para(doc, s.story_text)
        _para(doc, f"Requirements: {', '.join(s.requirement_ids)} · Priority: {s.priority.value if s.priority else '-'} · "
                   f"Review: {_status_label(s)}", size=8, color=RGBColor(0x59, 0x59, 0x59))
        for a in s.acceptance_criteria:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{a.ac_id} ").bold = True
            p.add_run(f"Given {a.given}, when {a.when}, then {a.then}.")
            if a.unverified_values:
                p.add_run(f"  [check: {', '.join(a.unverified_values)} not in sources]").font.color.rgb = RED

    doc.add_heading("8. Important test scenarios", level=1)
    if scenarios:
        _table(doc, ["ID", "Story", "Scenario", "Type / priority", "Steps", "Expected result"],
               [[t.scenario_id, t.story_id, t.title, f"{t.scenario_type.value}\n{t.priority.value}",
                 "\n".join(f"{i}. {x}" for i, x in enumerate(t.steps, start=1)), t.expected_result]
                for t in scenarios], [1.4, 1.4, 4.2, 2.2, 4.6, 4.2])
    else:
        _para(doc, "No test scenarios generated yet.", italic=True)

    doc.add_heading("9. Open items", level=1)
    doc.add_heading("Issues (contradictions, gaps, ambiguities)", level=2)
    if open_f:
        _table(doc, ["ID", "Type", "Severity", "Issue", "Requirements"],
               [[f.finding_id, FINDING_LABELS[f.finding_type], f.severity.value, f"{f.title}\n{f.description}",
                 ", ".join(f.requirement_ids) or "-"]
                for f in sorted(open_f, key=lambda f: SEVERITY_ORDER[f.severity])], [1.6, 2.6, 1.6, 9.0, 3.2])
    else:
        _para(doc, "No open issues.", italic=True)
    doc.add_heading("Clarification questions", level=2)
    questions = clarification_questions(b.findings, b.requirements, b.statements)
    if questions:
        _table(doc, ["#", "Ask", "Question", "Source"],
               [[str(i), q.ask_whom, q.question, q.source] for i, q in enumerate(questions, start=1)], [0.8, 3.8, 10.8, 2.6])
    else:
        _para(doc, "No open questions.", italic=True)
    not_accepted = [r for r in b.requirements if r.review_status not in EXCLUDED_REVIEW
                    and r.review_status.value not in ("APPROVED", "EDITED")]
    if b.final and not_accepted:
        doc.add_heading("Requirements not yet approved", level=2)
        _bullets(doc, [f"{r.requirement_id}: {r.title} ({_status_label(r)})" for r in not_accepted])

    doc.add_heading("10. Traceability matrix", level=1)
    rows = traceability(b.requirements, b.stories, b.scenarios, b.findings, b.docs)
    _table(doc, ["Requirement", "Review", "Statements", "Stories", "Tests", "Findings", "Coverage gaps"],
           [[t.requirement_id, t.review_status, t.statements, t.stories or "-", t.scenarios or "-", t.findings or "-",
             "; ".join(t.gaps) or "OK"] for t in rows], [2.0, 2.2, 2.6, 1.8, 2.6, 2.4, 4.4], 7)

    doc.add_heading("11. Glossary", level=1)
    terms = [(t, m or "Definition not provided") for t, m in b.glossary]
    terms += [(g.term, g.meaning or "Not defined in sources") for x in b.understandings.values() for g in x.glossary
              if g.term.lower() not in {t.lower() for t, _ in terms}]
    if terms:
        _table(doc, ["Term", "Meaning"], [[t, m] for t, m in terms], [5, 12], 9)
    else:
        _para(doc, NOT_FOUND, italic=True)

    doc.add_heading("12. Source register", level=1)
    _table(doc, ["ID", "File", "Type", "Segments", "Fingerprint"],
           [[d.document_id, d.file_name, DOCUMENT_TYPE_LABELS[d.document_type], str(d.stats.segments), d.sha256[:12]]
            for d in b.docs.values()], [1.6, 6.4, 4.0, 2.0, 3.0], 9)

    doc.add_heading("13. Review and approval", level=1)
    if b.audit:
        _table(doc, ["When (UTC)", "Item", "Action", "Change", "Reviewer", "Note"],
               [[a["ts"].replace("T", " ")[:16], a["item_id"], a["action"],
                 f"{a['field'] or ''}: {a['old_value'] or ''} → {a['new_value'] or ''}"[:160],
                 a["reviewer"] or "-", a["note"] or ""] for a in b.audit[-200:]], [2.6, 1.8, 2.2, 6.0, 2.2, 3.2], 7)
    else:
        _para(doc, "No review actions recorded yet.", italic=True)
    _table(doc, ["Role", "Name", "Signature", "Date"],
           [["Business owner", "", "", ""], ["Business analyst", "", "", ""], ["QA lead", "", "", ""]],
           [4, 5, 5, 3], 9)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- review report
def build_review_report(b: ExportBundle) -> bytes:
    """Short document for a clarification meeting: conflicts, issues by severity, questions by person."""
    doc = Document()
    _setup(doc, f"Requirement review – {b.session_name}", "DEMO DATA" if b.run.is_demo else None)
    t = doc.add_paragraph().add_run("Requirement review report")
    t.bold, t.font.size, t.font.color.rgb = True, Pt(22), BLUE
    _para(doc, f"{b.session_name} · run {b.run.run_id} · {datetime.now(timezone.utc):%d %b %Y}", color=RGBColor(0x59, 0x59, 0x59))
    findings = [f for f in b.findings if f.review_status != ReviewStatus.REJECTED]
    counts = defaultdict(int)
    for f in findings:
        counts[f.severity.value] += 1
    _para(doc, f"{len(findings)} issue(s): {counts['HIGH']} high, {counts['MEDIUM']} medium, {counts['LOW']} low. "
               f"{sum(1 for f in findings if f.review_status == ReviewStatus.RESOLVED)} resolved.")

    doc.add_heading("1. Contradictions between sources", level=1)
    conflicts = [f for f in findings if f.finding_type == FindingType.CONTRADICTION]
    if not conflicts:
        _para(doc, "No contradictions found.", italic=True)
    for f in conflicts:
        doc.add_heading(f"{f.finding_id} · {f.title}", level=3)
        if f.values:
            _table(doc, ["Requirement", "Value", "Source"],
                   [[v.requirement_id, v.value, ", ".join(b.docs[d].file_name if d in b.docs else d for d in v.document_ids)]
                    for v in f.values], [3, 4, 10], 9)
        else:
            _para(doc, f.description)
        _para(doc, f"Question: {f.clarification_question}", bold=True)
        if f.resolution:
            _para(doc, f"Resolution: {f.resolution}", color=RGBColor(0x00, 0x61, 0x00))

    comparison = compare_meeting_vs_brd(b.requirements, b.findings, b.docs)
    if comparison:
        doc.add_heading("2. Meeting compared with the BRD", level=1)
        _table(doc, ["Category", "Requirement", "Detail"],
               [[c.category, f"{c.requirement_id} {c.title}", c.detail] for c in comparison
                if c.category != "In both"], [3.2, 7.0, 7.4], 8)

    doc.add_heading("3. All issues by severity", level=1)
    _table(doc, ["ID", "Severity", "Type", "Issue", "Status"],
           [[f.finding_id, f.severity.value, FINDING_LABELS[f.finding_type], f"{f.title}\n{f.description}",
             f.review_status.value] for f in sorted(findings, key=lambda f: SEVERITY_ORDER[f.severity])],
           [1.6, 1.6, 2.8, 9.0, 2.4], 8)

    doc.add_heading("4. Questions to ask, by person", level=1)
    by_person: dict[str, list] = defaultdict(list)
    for q in clarification_questions(b.findings, b.requirements, b.statements):
        by_person[q.ask_whom].append(q)
    if not by_person:
        _para(doc, "No open questions.", italic=True)
    for person, qs in sorted(by_person.items()):
        doc.add_heading(person, level=3)
        for q in qs:
            doc.add_paragraph(f"{q.question}  ({q.source})", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
